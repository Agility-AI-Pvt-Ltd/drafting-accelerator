"""
Generic parser: reads counts and capacities out of whatever mass-balance and
design-criteria files are uploaded (same sheet/row template as the NDDB dummy
data), instead of hardcoding equipment lists for one fixed file.
"""
import re
import openpyxl

WORD_NUM = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7, "eight": 8}

TYPE_TAG_PREFIX = {"reception": "TK", "despatch": "TD", "separator": "SEP", "processing": "PL"}
TYPE_LABEL = {"reception": "Tanker reception line", "despatch": "Tanker despatch line",
              "separator": "Separator", "processing": "Milk processing line"}
TYPE_X = {"reception": 0, "despatch": 60, "separator": 95, "processing": 150}
BOX_W, BOX_H = 40, 14


def _capacity(text):
    if not text:
        return None
    m = re.search(r"(\d+)\s*KLPH", text, re.I)
    return int(m.group(1)) if m else None


def _leading_count(text):
    if not text:
        return None
    m = re.match(r"\s*(One|Two|Three|Four|Five|Six|Seven|Eight)\b", text, re.I)
    return WORD_NUM.get(m.group(1).lower()) if m else None


def _reception_counts(desc_text, note_text):
    base = expanded = None
    if note_text:
        m = re.search(r"(\d+)\s*nos", note_text, re.I)
        if m:
            base = int(m.group(1))
        m2 = re.search(r"to\s*(\d+)", note_text, re.I)
        if m2:
            expanded = int(m2.group(1))
    if base is None:
        base = _leading_count(desc_text)
    if expanded is None:
        expanded = base
    return base, expanded


def _reception_spec(desc_text, capacity):
    m = re.search(r"\((.*?)\)", desc_text or "")
    config = m.group(1).replace(" ", "") if m else ""
    parts = [p for p in [config, f"{capacity} KLPH" if capacity else None] if p]
    return ", ".join(parts)


def _separator_spec(desc_text, capacity):
    bits = [f"{capacity} KLPH" if capacity else None]
    if desc_text and "self" in desc_text.lower():
        bits.append("self-cleaning")
    if desc_text and "fat" in desc_text.lower():
        bits.append("auto fat std.")
    return ", ".join(b for b in bits if b)


def parse_design_criteria(path):
    """Reads the 'Design criteria' sheet and returns counts/specs per equipment type."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Design criteria"]

    title = ws["A1"].value
    reception_desc, reception_note = ws["B8"].value, ws["D8"].value
    despatch_desc = ws["B9"].value
    processing_desc = ws["B10"].value
    separator_desc = ws["B11"].value

    reception_base, reception_expanded = _reception_counts(reception_desc, reception_note)
    despatch_count = _leading_count(despatch_desc) or 1
    processing_count = _leading_count(processing_desc) or 1
    separator_count = _leading_count(separator_desc) or 1

    reception_cap = _capacity(reception_desc)
    despatch_cap = _capacity(despatch_desc)
    processing_cap = _capacity(processing_desc)
    separator_cap = _capacity(separator_desc)

    return {
        "title": title,
        "source_note": f"{reception_desc}  |  {reception_note}",
        "reception_base": reception_base,
        "reception_expanded": reception_expanded,
        "counts": {
            "reception": reception_base,
            "despatch": despatch_count,
            "separator": separator_count,
            "processing": processing_count,
        },
        "counts_expanded": {
            "reception": reception_expanded,
            "despatch": despatch_count,
            "separator": separator_count,
            "processing": processing_count,
        },
        "specs": {
            "reception": _reception_spec(reception_desc, reception_cap),
            "despatch": f"{despatch_cap} KLPH" if despatch_cap else "",
            "separator": _separator_spec(separator_desc, separator_cap),
            "processing": f"{processing_cap} KLPH" if processing_cap else "",
        },
    }


def build_equipment_list(counts, specs):
    """Turns {type: count} + {type: spec text} into tagged, laid-out equipment boxes."""
    equipment = []
    for etype, count in counts.items():
        if not count:
            continue
        for i in range(count):
            tag = f"{TYPE_TAG_PREFIX[etype]}-{101 + i}"
            equipment.append({
                "tag": tag,
                "name": f"{TYPE_LABEL[etype]} {i + 1}",
                "spec": specs.get(etype, ""),
                "type": etype,
                "x": TYPE_X[etype],
                "y": i * 20,
                "w": BOX_W,
                "h": BOX_H,
            })
    return equipment


def diff_equipment(list_a, list_b):
    """Marks each item in list_b as changed if it's new or its spec differs from list_a."""
    by_tag_a = {e["tag"]: e for e in list_a}
    out = []
    for e in list_b:
        prev = by_tag_a.get(e["tag"])
        changed = prev is None or prev["spec"] != e["spec"]
        out.append({**e, "changed": changed})
    return out


def extract_mass_balance(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Sheet1"]
    products = []
    for row in range(4, 17):
        name = ws[f"H{row}"].value
        if not name:
            continue
        products.append({"name": name, "kg": ws[f"J{row}"].value})
    return {
        "milk_in_l": ws["B4"].value,
        "milk_in_kg": ws["C4"].value,
        "fat_pct": ws["D4"].value,
        "snf_pct": ws["E4"].value,
        "kg_fat": ws["F4"].value,
        "kg_snf": ws["G4"].value,
        "products": products,
    }
